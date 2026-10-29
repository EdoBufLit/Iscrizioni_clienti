from datetime import datetime, timedelta, timezone
from io import BytesIO
import uuid

import pytest
from openpyxl import load_workbook

from app.db import SessionLocal
from app.models import AdminRole, AdminUser, Member, MemberDocument, MemberStatus, OrgAdminToken, Organization
from app.services.member_filters import build_org_member_query
from app.utils import hash_token


@pytest.fixture
def tenant(client):
    db = SessionLocal()
    key = uuid.uuid4().hex
    org = Organization(name="Mobile QA", slug=f"mobile-{key}", is_active=True)
    db.add(org)
    db.flush()
    admin = AdminUser(email=f"mobile-{key}@example.com", org_id=org.id, role=AdminRole.ORG_ADMIN, is_active=True)
    db.add(admin)
    db.flush()
    db.add(OrgAdminToken(admin_id=admin.id, token_hash=hash_token(key), expires_at=datetime.utcnow()+timedelta(minutes=10)))
    db.commit()
    assert client.get(f"/api/org-admin/auth/verify?token={key}", follow_redirects=False).status_code in (302, 303)
    yield db, org
    db.close()


def add_member(db, org, **overrides):
    number = 100000 + db.query(Member).filter(Member.org_id == org.id).count()
    values = dict(org_id=org.id, first_name="Socio", status=MemberStatus.ACTIVE, card_no=number,
                  card_year=datetime.utcnow().year, valid_until=datetime.utcnow()+timedelta(days=60))
    values.update(overrides)
    member = Member(**values)
    db.add(member)
    db.flush()
    return member


def test_expiring_list_and_global_metric_exclude_expired_deleted_and_other_tenants(client, tenant):
    db, org = tenant
    due = add_member(db, org, valid_until=datetime.utcnow()+timedelta(days=10))
    add_member(db, org)
    add_member(db, org, status=MemberStatus.EXPIRED, valid_until=datetime.utcnow()-timedelta(days=2))
    add_member(db, org, deleted_at=datetime.utcnow(), valid_until=datetime.utcnow()+timedelta(days=1))
    other = Organization(name="Other", slug=f"other-{uuid.uuid4().hex}", is_active=True)
    db.add(other); db.flush()
    add_member(db, other, valid_until=datetime.utcnow()+timedelta(days=1))
    db.commit()
    result = client.get("/api/org-admin/members?status=expiring").json()
    assert result["total"] == 1
    assert [item["id"] for item in result["items"]] == [due.id]
    assert client.get("/api/org-admin/metrics").json()["expiring_members_count"] == 1


def test_review_counts_members_once_and_bulk_uses_the_same_filter(client, tenant):
    db, org = tenant
    pending = add_member(db, org, status=MemberStatus.PENDING_VERIFICATION, card_no=None)
    docs = add_member(db, org)
    deleted = add_member(db, org, deleted_at=datetime.utcnow())
    rejected = add_member(db, org, status=MemberStatus.REJECTED)
    expired = add_member(db, org, status=MemberStatus.EXPIRED)
    other = Organization(name="Other review", slug=f"review-{uuid.uuid4().hex}", is_active=True)
    db.add(other); db.flush()
    foreign = add_member(db, other, status=MemberStatus.PENDING_VERIFICATION, card_no=None)
    for member in (docs, deleted, rejected, expired, foreign):
        for kind in ("identity", "fiscal_code"):
            db.add(MemberDocument(member_id=member.id, doc_type=kind, status="pending" if kind == "identity" else "uploaded"))
    db.commit()
    result = client.get("/api/org-admin/members?status=review").json()
    assert result["total"] == 2
    assert {item["id"] for item in result["items"]} == {pending.id, docs.id}
    assert client.get("/api/org-admin/metrics").json()["members_to_review_count"] == 2
    run = client.post("/api/org-admin/member-bulk/preview", json={"request_key": uuid.uuid4().hex,
        "action": "export", "selection": {"mode": "filtered", "filters": {"status": "review"}}})
    assert run.status_code == 200, run.text
    assert run.json()["selected_count"] == 2
    assert {item["member_id"] for item in run.json()["items"]} == {pending.id, docs.id}


def test_expiring_legacy_annual_card_uses_actual_rome_expiry(tenant):
    db, org = tenant
    member = add_member(db, org, card_year=2026, valid_until=None)
    db.commit()
    assert build_org_member_query(db, org_id=org.id, status="expiring", now=datetime(2026,12,15)).one().id == member.id
    assert build_org_member_query(db, org_id=org.id, status="expiring", now=datetime(2026,10,28)).count() == 0
    assert build_org_member_query(db, org_id=org.id, status="expiring", now=datetime(2026,12,15,tzinfo=timezone.utc)).one().id == member.id


def test_pending_documents_metric_matches_filtered_members_not_document_count(client, tenant):
    db, org = tenant
    active = add_member(db, org)
    pending = add_member(db, org, status=MemberStatus.PENDING_DOCS)
    pending_without_documents = add_member(db, org, status=MemberStatus.PENDING_VERIFICATION, card_no=None)
    excluded = [
        add_member(db, org, status=MemberStatus.EXPIRED),
        add_member(db, org, status=MemberStatus.REJECTED),
        add_member(db, org, deleted_at=datetime.utcnow()),
    ]
    other = Organization(name="Other documents", slug=f"other-docs-{uuid.uuid4().hex}", is_active=True)
    db.add(other); db.flush()
    excluded.append(add_member(db, other))
    for member in [active, *excluded]:
        for kind, status in (("identity", "pending"), ("fiscal_code", "uploaded")):
            db.add(MemberDocument(member_id=member.id, doc_type=kind, status=status))
    db.add(MemberDocument(member_id=pending.id, doc_type="identity", status="uploaded"))
    db.commit()

    response = client.get("/api/org-admin/members", params={"status": "review", "docs": "pending"})
    assert response.status_code == 200, response.text
    filtered = response.json()
    assert filtered["total"] == 2
    assert {item["id"] for item in filtered["items"]} == {active.id, pending.id}
    metrics = client.get("/api/org-admin/metrics").json()
    assert metrics["members_pending_documents_review_count"] == filtered["total"]
    assert metrics["documents_pending_review"] == 3  # Existing document-based metric is unchanged.
    assert metrics["pending_requests_count"] == 2
    assert pending_without_documents.id not in {item["id"] for item in filtered["items"]}
    pending_list = client.get("/api/org-admin/members", params={"status": "pending"}).json()
    assert pending_list["total"] == metrics["pending_requests_count"]
    assert {item["id"] for item in pending_list["items"]} == {pending.id, pending_without_documents.id}


def workbook_members(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    rows = list(workbook.active.iter_rows(min_row=5, values_only=True))
    workbook.close()
    return [row for row in rows if row[0] is not None]


def test_single_member_xlsx_contains_only_requested_member_and_global_export_is_unchanged(client, tenant):
    db, org = tenant
    first = add_member(db, org, first_name="Primo", last_name="Socio")
    second = add_member(db, org, first_name="Secondo", last_name="Socio")
    expired = add_member(db, org, first_name="Scaduto", last_name="Consultabile", status=MemberStatus.EXPIRED)
    db.commit()

    response = client.get("/api/org-admin/members.xlsx", params={"member_id": second.id})
    assert f"scheda-socio-{second.id}-" in response.headers["content-disposition"]
    rows = workbook_members(response)
    assert len(rows) == 1 and rows[0][0] == "Secondo"
    assert rows[0][6] == second.card_no

    # A visible expired profile can be exported individually, while the global
    # register keeps its existing active/pending eligibility rules.
    assert workbook_members(client.get("/api/org-admin/members.xlsx", params={"member_id": expired.id}))[0][0] == "Scaduto"
    global_response = client.get("/api/org-admin/members.xlsx")
    assert "libro-soci-" in global_response.headers["content-disposition"]
    assert {row[0] for row in workbook_members(global_response)} == {first.first_name, second.first_name}


def test_single_member_xlsx_rejects_missing_foreign_deleted_and_unauthenticated_requests(client, tenant):
    db, org = tenant
    own = add_member(db, org)
    deleted = add_member(db, org, deleted_at=datetime.utcnow())
    other = Organization(name="Foreign export", slug=f"foreign-export-{uuid.uuid4().hex}", is_active=True)
    db.add(other); db.flush()
    foreign = add_member(db, other)
    db.commit()
    for member_id in (deleted.id, foreign.id, 999999999):
        response = client.get("/api/org-admin/members.xlsx", params={"member_id": member_id})
        assert response.status_code == 404, response.text
    client.cookies.clear()
    assert client.get("/api/org-admin/members.xlsx", params={"member_id": own.id}).status_code == 401


@pytest.mark.parametrize("member_id", ["0", "-1", "1.5", "abc"])
def test_single_member_xlsx_requires_a_positive_integer(client, tenant, member_id):
    assert client.get("/api/org-admin/members.xlsx", params={"member_id": member_id}).status_code == 422
