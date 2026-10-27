from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
import uuid

import pytest
from openpyxl import load_workbook
from sqlalchemy import event

from app.db import Base, SessionLocal, engine
from app.main import app
from app.models import AdminMfaFactor, AdminRole, AdminUser, AnnualMembershipTerm, EmailOutbox, Member, MemberDocument, MemberNotification, MemberStatus, OrgAdminSession, OrgAdminToken, Organization, Token
from app.models_member_bulk import MemberBulkItem, MemberBulkRun
from app import models_accounting as _models_accounting  # noqa: F401 -- include read-model tables
from app.routes.org_admin_member_bulk import router
from app.services import member_bulk
from app.utils import hash_token


PREFIX = "/api/org-admin/member-bulk"


@pytest.fixture
def tenant(client):
    if not any(getattr(route, "path", None) == f"{PREFIX}/preview" for route in app.routes):
        app.include_router(router)
    Base.metadata.create_all(engine)
    db = SessionLocal()
    suffix = uuid.uuid4().hex[:10]
    org = Organization(name="Bulk locale", slug=f"bulk-{suffix}", is_active=True, membership_fee_amount=Decimal("30.00"), membership_fee_currency="EUR")
    db.add(org)
    db.flush()
    admin = AdminUser(email=f"admin-{suffix}@example.com", role=AdminRole.ORG_ADMIN, org_id=org.id, is_active=True)
    db.add(admin)
    db.flush()
    login(client, db, admin)
    yield db, org, admin
    db.close()


def login(client, db, admin):
    raw = uuid.uuid4().hex
    db.add(OrgAdminToken(admin_id=admin.id, token_hash=hash_token(raw), expires_at=datetime.utcnow() + timedelta(minutes=10)))
    db.commit()
    response = client.get(f"/api/org-admin/auth/verify?token={raw}", follow_redirects=False)
    assert response.status_code in {302, 303}


def make_members(db, org, count=1, **overrides):
    suffix = uuid.uuid4().hex[:9]
    members = []
    for index in range(count):
        fields = dict(org_id=org.id, first_name="Socio", last_name=f"Bulk {index:03}", email=f"{suffix}-{index}@example.com", status=MemberStatus.ACTIVE, card_no=700000 + index, card_year=datetime.utcnow().year, membership_type="annual", valid_from=datetime.utcnow() - timedelta(days=1), valid_until=datetime.utcnow() + timedelta(days=350), joined_at=datetime.utcnow(), membership_fee_snapshot=Decimal("30.00"))
        fields.update(overrides)
        members.append(Member(**fields))
    db.add_all(members)
    db.commit()
    return members


def preview(client, members=(), *, action="send_access", selection=None, year=None, key=None):
    body = dict(request_key=key or uuid.uuid4().hex, action=action, selection=selection or {"mode": "ids", "member_ids": [member.id for member in members]})
    if year is not None:
        body["membership_year"] = year
    response = client.post(f"{PREFIX}/preview", json=body)
    assert response.status_code == 200, response.text
    return response.json()


def execute(client, run, *, key=None):
    response = client.post(f"{PREFIX}/{run['id']}/execute", json={"request_key": key or uuid.uuid4().hex})
    assert response.status_code == 200, response.text
    return response.json()


def freeze_bulk_clock(monkeypatch, now):
    real_now = member_bulk._now
    monkeypatch.setattr(member_bulk, "_now", lambda value=None: real_now(value if value is not None else now))


def test_preview_freezes_selection_and_chunks_are_idempotent(client, tenant):
    db, org, _admin = tenant
    members = make_members(db, org, 31)
    run = preview(client, members)
    assert run["selected_count"] == run["eligible_count"] == 31
    assert run["status"] == "preview" and run["counts"]["pending"] == 31
    assert db.query(EmailOutbox).count() == db.query(Token).filter(Token.member_id.in_([m.id for m in members])).count() == 0
    key = uuid.uuid4().hex
    first = execute(client, run, key=key)
    assert first["counts"] == dict(pending=6, succeeded=25, skipped=0, failed=0)
    assert first["status"] == "running"
    repeated = execute(client, run, key=key)
    assert repeated["counts"] == first["counts"]
    completed = execute(client, run)
    assert completed["status"] == "completed" and completed["counts"]["succeeded"] == 31
    assert db.query(EmailOutbox).count() == 31
    assert db.query(Token).filter(Token.member_id.in_([m.id for m in members])).count() == 31
    page = client.get(f"{PREFIX}/{run['id']}?offset=25&limit=5").json()
    assert len(page["items"]) == 5 and page["has_more"] and page["total"] == 31
    assert {item["delivery_status"] for item in page["items"]} == {"queued"}
    recent = client.get(f"{PREFIX}?limit=1").json()
    assert recent["items"][0]["id"] == run["id"] and recent["items"][0]["counts"]["succeeded"] == 31
    assert "items" not in recent["items"][0] and "name" not in str(recent["items"][0])


def test_preview_request_key_cannot_change_payload(client, tenant):
    db, org, _admin = tenant
    member = make_members(db, org)[0]
    key = uuid.uuid4().hex
    first = preview(client, [member], key=key)
    second = preview(client, [member], key=key)
    assert first["id"] == second["id"]
    response = client.post(f"{PREFIX}/preview", json={"request_key": key, "action": "export", "selection": {"mode": "ids", "member_ids": [member.id]}})
    assert response.status_code == 409
    assert db.query(MemberBulkRun).filter(MemberBulkRun.org_id == org.id).count() == 1


def test_all_filtered_selection_matches_exact_and_fuzzy_list(client, tenant):
    db, org, _admin = tenant
    exact, fuzzy, other = make_members(db, org, 3)
    exact.first_name, exact.last_name = "Mario", "Rossi"
    fuzzy.first_name, fuzzy.last_name = "Mario", "Rossu"
    other.first_name, other.last_name = "Elena", "Bianchi"
    for member in (exact, fuzzy, other):
        member.is_manual = True
    db.add_all([MemberDocument(member_id=member.id, doc_type="identity", status="approved") for member in (exact, fuzzy)])
    db.commit()
    filters = dict(q="Mario Rossi", status="active", source="manual", docs="approved", order="last_name_asc")
    listed = client.get("/api/org-admin/members", params={**filters, "limit": 100}).json()
    listed_ids = [item["id"] for item in listed["items"]]
    assert set(listed_ids) == {exact.id, fuzzy.id}
    run = preview(client, action="export", selection={"mode": "filtered", "filters": filters})
    assert [item["member_id"] for item in run["items"]] == listed_ids
    excluded = preview(client, action="export", selection={"mode": "filtered", "filters": filters, "excluded_ids": [exact.id]})
    assert [item["member_id"] for item in excluded["items"]] == [fuzzy.id]


def test_selection_cap_is_explicit_never_truncated(client, tenant, monkeypatch):
    db, org, _admin = tenant
    make_members(db, org)
    monkeypatch.setattr(member_bulk, "MAX_SELECTION", 3)
    monkeypatch.setattr(member_bulk, "iter_member_search_ids", lambda *_: iter(range(100, 110)))
    response = client.post(f"{PREFIX}/preview", json={"request_key": uuid.uuid4().hex, "action": "export", "selection": {"mode": "filtered", "filters": {"status": "all"}}})
    assert response.status_code == 422 and "10.000" in response.text
    assert db.query(MemberBulkRun).filter(MemberBulkRun.org_id == org.id).count() == 0


def test_foreign_member_run_and_other_actor_are_rejected(client, tenant):
    db, org, admin = tenant
    own = make_members(db, org)[0]
    foreign = Organization(name="Altra", slug=f"foreign-{uuid.uuid4().hex}", is_active=True)
    db.add(foreign)
    db.commit()
    outsider = make_members(db, foreign)[0]
    response = client.post(f"{PREFIX}/preview", json={"request_key": uuid.uuid4().hex, "action": "export", "selection": {"mode": "ids", "member_ids": [own.id, outsider.id]}})
    assert response.status_code == 404 and outsider.email not in response.text
    run = preview(client, [own], action="export")
    other_admin = AdminUser(email=f"other-{uuid.uuid4().hex}@example.com", org_id=org.id, role=AdminRole.ORG_ADMIN, is_active=True)
    db.add(other_admin)
    db.flush()
    login(client, db, other_admin)
    assert client.get(f"{PREFIX}/{run['id']}").status_code == 404
    assert client.post(f"{PREFIX}/{run['id']}/execute", json={"request_key": uuid.uuid4().hex}).status_code == 404
    assert client.get(PREFIX).json()["total"] == 0
    login(client, db, admin)
    org.is_active = False
    db.commit()
    assert client.get(f"{PREFIX}/{run['id']}").status_code == 403
    assert client.get(PREFIX).status_code == 403
    assert client.post(f"{PREFIX}/{run['id']}/execute", json={"request_key": uuid.uuid4().hex}).status_code == 403
    assert db.query(EmailOutbox).count() == 0


def test_changed_recipient_and_fresh_eligibility_are_skipped(client, tenant):
    db, org, _admin = tenant
    changed, activated, removed = make_members(db, org, 3)
    run = preview(client, [changed, activated, removed])
    changed.email = "changed-address@example.com"
    activated.password_hash = "a-password-was-configured"
    removed.deleted_at = datetime.utcnow()
    db.commit()
    result = execute(client, run)
    assert result["counts"] == dict(pending=0, succeeded=0, skipped=3, failed=0)
    assert set(result["reason_counts"]) == {"recipient_changed", "access_already_active", "member_unavailable"}
    assert db.query(EmailOutbox).count() == 0


def test_no_eligible_members_and_throttle_across_runs(client, tenant):
    db, org, _admin = tenant
    active, configured, invalid, pending = make_members(db, org, 4)
    configured.password_hash = "present"
    invalid.email = None
    pending.status = MemberStatus.PENDING_DOCS
    db.commit()
    skipped = preview(client, [configured, invalid, pending])
    assert skipped["eligible_count"] == 0 and not skipped["can_execute"]
    assert skipped["counts"]["skipped"] == 3
    first = execute(client, preview(client, [active]))
    assert first["counts"]["succeeded"] == 1
    second = preview(client, [active])
    assert second["eligible_count"] == 0 and second["reason_counts"] == {"access_recently_queued": 1}
    legacy = client.post(f"/api/org-admin/members/{active.id}/send-access")
    assert legacy.status_code == 429
    assert db.query(EmailOutbox).count() == 1


def test_only_failed_rows_are_retried_with_atomic_rollback(client, tenant, monkeypatch):
    db, org, _admin = tenant
    good, failing, skipped = make_members(db, org, 3)
    skipped.password_hash = "configured"
    db.commit()
    run = preview(client, [good, failing, skipped])
    queue = member_bulk._queue_access

    def failing_after_queue(*args, **kwargs):
        result = queue(*args, **kwargs)
        if kwargs["member"].id == failing.id:
            raise RuntimeError("synthetic failure after queue, before transaction commit")
        return result

    monkeypatch.setattr(member_bulk, "_queue_access", failing_after_queue)
    first = execute(client, run)
    assert first["counts"] == dict(pending=0, succeeded=1, skipped=1, failed=1)
    assert db.query(EmailOutbox).count() == 1
    assert db.query(Token).filter(Token.member_id == failing.id).count() == 0
    key = uuid.uuid4().hex
    retry = client.post(f"{PREFIX}/{run['id']}/retry-failed", json={"request_key": key})
    assert retry.status_code == 200 and retry.json()["counts"]["pending"] == 1
    failed_again = execute(client, run)
    assert failed_again["counts"]["failed"] == 1
    repeated_retry = client.post(f"{PREFIX}/{run['id']}/retry-failed", json={"request_key": key})
    assert repeated_retry.json()["counts"]["failed"] == 1 and repeated_retry.json()["counts"]["pending"] == 0
    monkeypatch.setattr(member_bulk, "_queue_access", queue)
    client.post(f"{PREFIX}/{run['id']}/retry-failed", json={"request_key": uuid.uuid4().hex})
    complete = execute(client, run)
    assert complete["counts"] == dict(pending=0, succeeded=2, skipped=1, failed=0)
    assert db.query(EmailOutbox).count() == 2
    assert {item["member_id"]: item["attempts"] for item in complete["items"]} == {good.id: 1, failing.id: 3, skipped.id: 0}


def test_preview_expires_but_started_runs_remain_resumable(client, tenant):
    db, org, _admin = tenant
    member = make_members(db, org)[0]
    previewed = preview(client, [member], action="export")
    db.query(MemberBulkRun).filter_by(id=previewed["id"]).update({"expires_at": datetime.now(timezone.utc) - timedelta(seconds=1)})
    db.commit()
    assert client.get(f"{PREFIX}/{previewed['id']}").json()["status"] == "expired"
    assert client.post(f"{PREFIX}/{previewed['id']}/execute", json={"request_key": uuid.uuid4().hex}).status_code == 409
    assert client.get(f"{PREFIX}/{previewed['id']}/export.xlsx").status_code == 409
    running = preview(client, [member], action="export")
    db.query(MemberBulkRun).filter_by(id=running["id"]).update({"status": "running", "started_at": datetime.now(timezone.utc) - timedelta(hours=2), "expires_at": datetime.now(timezone.utc) - timedelta(hours=1)})
    db.commit()
    assert execute(client, running)["status"] == "completed"


def test_export_is_confirmed_selected_and_formula_safe(client, tenant):
    db, org, _admin = tenant
    chosen, omitted = make_members(db, org, 2)
    chosen.first_name = '=HYPERLINK("https://example.invalid","test")'
    chosen.status = MemberStatus.EXPIRED
    db.commit()
    run = preview(client, [chosen], action="export")
    assert client.get(f"{PREFIX}/{run['id']}/export.xlsx").status_code == 409
    execute(client, run)
    response = client.get(f"{PREFIX}/{run['id']}/export.xlsx")
    assert response.status_code == 200 and response.headers["x-exported-count"] == "1"
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    assert sheet["A5"].data_type != "f" and sheet["A5"].value.startswith("'=")
    assert sheet["C5"].value == chosen.email and omitted.email not in response.text


@pytest.mark.parametrize("bulk_delete", [False, True])
def test_purged_integer_id_never_targets_replacement_profile(client, tenant, bulk_delete):
    db, org, _admin = tenant
    member = make_members(db, org)[0]
    requested_id = member.id
    org_id = org.id
    run = preview(client, [member], action="export")
    if bulk_delete:
        db.query(Member).filter(Member.id == requested_id).delete(synchronize_session=False)
    else:
        db.delete(member)
    db.commit()
    db.expunge_all()
    replacement = Member(id=requested_id, org_id=org_id, first_name="Replacement", last_name="Do not export", email="replacement@example.com", status=MemberStatus.ACTIVE)
    db.add(replacement)
    db.commit()
    item = db.query(MemberBulkItem).filter_by(run_id=run["id"]).one()
    assert item.member_id is None and item.requested_member_id == requested_id
    result = execute(client, run)
    assert result["counts"]["skipped"] == 1 and result["items"][0]["member_id"] == requested_id
    response = client.get(f"{PREFIX}/{run['id']}/export.xlsx")
    assert response.status_code == 200 and response.headers["x-exported-count"] == "0"
    assert "Replacement" not in str(list(load_workbook(BytesIO(response.content)).active.values))


def test_quote_reminders_use_real_balance_and_skip_new_payments(client, tenant, monkeypatch):
    db, org, _admin = tenant
    freeze_bulk_clock(monkeypatch, datetime(2027, 2, 1, tzinfo=timezone.utc))
    year = 2027
    due, paid, changed = make_members(db, org, 3, card_year=year)
    balances = {due.id: 1800, paid.id: 0, changed.id: 3000}
    import app.services.member_accounting as accounting

    def read_balances(_db, *, members, year):
        return {member.id: dict(expected_cents=3000, paid_cents=3000 - balances[member.id], outstanding_cents=balances[member.id], currency="EUR") for member in members}

    monkeypatch.setattr(accounting, "member_year_balances", read_balances)
    run = preview(client, [due, paid, changed], action="quote_reminder", year=year)
    assert run["eligible_count"] == 2 and run["outstanding_total_cents"] == 4800
    balances[changed.id] = 1500
    result = execute(client, run)
    assert result["counts"] == dict(pending=0, succeeded=1, skipped=2, failed=0)
    assert result["outstanding_total_cents"] == 4800
    assert result["reason_counts"] == {"already_paid": 1, "balance_changed": 1}
    message = db.query(EmailOutbox).one()
    assert message.to_email == due.email and "18,00 EUR" in message.payload_json["text_body"]
    assert message.email_type == "membership_fee_reminder"
    assert db.query(MemberNotification).filter(MemberNotification.member_id == due.id).count() == 1
    repeated = preview(client, [due], action="quote_reminder", year=year)
    assert repeated["reason_counts"] == {"reminder_recently_queued": 1}


def test_quote_reminder_reads_accounting_projection_without_fabricating_debt(client, tenant, monkeypatch):
    db, org, _admin = tenant
    freeze_bulk_clock(monkeypatch, datetime(2027, 2, 1, tzinfo=timezone.utc))
    year = 2027
    due, free = make_members(db, org, 2, card_year=year)
    db.add_all([AnnualMembershipTerm(member_id=member.id, org_id=org.id, membership_year=year, starts_on=date(year, 1, 1), valid_through=date(year, 12, 31), status="due", fee_amount=fee, currency="EUR") for member, fee in [(due, Decimal("12.50")), (free, Decimal("0.00"))]])
    db.commit()
    run = preview(client, [due, free], action="quote_reminder", year=year)
    assert run["eligible_count"] == 1 and run["outstanding_total_cents"] == 1250
    assert run["reason_counts"] == {"no_fee_due": 1}
    result = execute(client, run)
    assert result["counts"]["succeeded"] == 1
    assert "12,50 EUR" in db.query(EmailOutbox).one().payload_json["text_body"]


@pytest.mark.parametrize("now,year", [
    (datetime(2026, 10, 27, tzinfo=timezone.utc), 2027),
    (datetime(2026, 12, 31, 22, 59, tzinfo=timezone.utc), 2027),
    (datetime(2027, 2, 1, tzinfo=timezone.utc), 2028),
    (datetime(2027, 2, 1, tzinfo=timezone.utc), 2026),
])
def test_quote_reminder_preview_rejects_unstarted_or_legacy_years(client, tenant, monkeypatch, now, year):
    db, org, _admin = tenant
    freeze_bulk_clock(monkeypatch, now)
    member = make_members(db, org)[0]
    response = client.post(f"{PREFIX}/preview", json={"request_key": uuid.uuid4().hex,
        "action": "quote_reminder", "membership_year": year,
        "selection": {"mode": "ids", "member_ids": [member.id]}})
    assert response.status_code == 409, response.text
    assert db.query(MemberBulkRun).filter_by(org_id=org.id).count() == 0
    assert db.query(EmailOutbox).filter_by(to_email=member.email).count() == 0


@pytest.mark.parametrize("now,year", [
    (datetime(2026, 10, 27, tzinfo=timezone.utc), 2027),
    (datetime(2027, 2, 1, tzinfo=timezone.utc), 2028),
])
def test_old_reminder_preview_cannot_be_resumed_or_retried_early(client, tenant, monkeypatch, now, year):
    db, org, _admin = tenant
    freeze_bulk_clock(monkeypatch, now)
    member = make_members(db, org)[0]
    legacy = preview(client, [member], action="export")
    run = db.get(MemberBulkRun, legacy["id"])
    run.action, run.membership_year = "quote_reminder", year
    db.commit()
    response = client.post(f"{PREFIX}/{run.id}/execute", json={"request_key": uuid.uuid4().hex})
    assert response.status_code == 409
    run.status = "completed"
    item = db.query(MemberBulkItem).filter_by(run_id=run.id).one()
    item.status = "failed"
    db.commit()
    response = client.post(f"{PREFIX}/{run.id}/retry-failed", json={"request_key": uuid.uuid4().hex})
    assert response.status_code == 409
    state = client.get(f"{PREFIX}/{run.id}").json()
    assert state["can_execute"] is False and state["can_retry_failed"] is False
    assert state["unavailable_reason"]
    assert db.query(EmailOutbox).filter_by(to_email=member.email).count() == 0


def test_quote_reminder_respects_previous_card_validity_on_rome_new_year(client, tenant, monkeypatch):
    db, org, _admin = tenant
    member = make_members(db, org, card_year=2026, valid_until=None)[0]
    db.add(AnnualMembershipTerm(member_id=member.id, org_id=org.id, membership_year=2026,
        starts_on=date(2026, 1, 1), valid_through=date(2027, 1, 1), card_no=member.card_no,
        status="active", fee_amount=Decimal("30.00"), currency="EUR"))
    db.add(AnnualMembershipTerm(member_id=member.id, org_id=org.id, membership_year=2027,
        starts_on=date(2027, 1, 1), valid_through=date(2028, 1, 1), status="due",
        fee_amount=Decimal("30.00"), currency="EUR"))
    db.commit()
    # UTC is still in 2026, but Rome has entered the new year. The prior card
    # nevertheless covers the whole of January 1, so no reminder is eligible.
    freeze_bulk_clock(monkeypatch, datetime(2026, 12, 31, 23, 0, tzinfo=timezone.utc))
    run = preview(client, [member], action="quote_reminder", year=2027)
    assert run["eligible_count"] == 0 and run["reason_counts"] == {"current_card_valid": 1}
    freeze_bulk_clock(monkeypatch, datetime(2027, 1, 1, 23, 0, tzinfo=timezone.utc))
    run = preview(client, [member], action="quote_reminder", year=2027)
    assert run["eligible_count"] == 1
    result = execute(client, run)
    assert result["counts"]["succeeded"] == 1


@pytest.mark.parametrize("member_id", [True, 1.0, "1", 0, -1])
def test_selection_rejects_coerced_or_nonpositive_ids(client, tenant, member_id):
    response = client.post(f"{PREFIX}/preview", json={"request_key": uuid.uuid4().hex, "action": "export", "selection": {"mode": "ids", "member_ids": [member_id]}})
    assert response.status_code == 422


def test_access_preview_queries_are_batched_not_per_member(client, tenant):
    db, org, _admin = tenant
    members = make_members(db, org, 70)
    selection = {"mode": "ids", "member_ids": [member.id for member in members]}
    statements = []

    def capture(_connection, _cursor, statement, _params, _context, _many):
        if statement.lstrip().upper().startswith("SELECT"):
            statements.append(statement)

    event.listen(engine, "before_cursor_execute", capture)
    try:
        run = preview(client, selection=selection)
    finally:
        event.remove(engine, "before_cursor_execute", capture)
    assert run["eligible_count"] == 70
    assert len(statements) <= 18, len(statements)


@pytest.mark.parametrize("year", [True, 2026.0, "2026", 1999, 2201])
def test_year_is_a_bounded_strict_integer(client, tenant, year):
    db, org, _admin = tenant
    member = make_members(db, org)[0]
    response = client.post(f"{PREFIX}/preview", json={"request_key": uuid.uuid4().hex, "action": "quote_reminder", "membership_year": year, "selection": {"mode": "ids", "member_ids": [member.id]}})
    assert response.status_code == 422


def test_new_mfa_requirement_invalidates_old_session(client, tenant):
    db, org, admin = tenant
    member = make_members(db, org)[0]
    run = preview(client, [member])
    db.add(AdminMfaFactor(admin_id=admin.id, secret_encrypted="synthetic-never-used", status="active"))
    db.commit()
    assert client.get(f"{PREFIX}/{run['id']}").status_code == 401
    assert client.post(f"{PREFIX}/{run['id']}/execute", json={"request_key": uuid.uuid4().hex}).status_code == 401
    assert db.query(EmailOutbox).count() == 0


def test_session_revocation_during_batch_aborts_uncommitted_messages(client, tenant, monkeypatch):
    db, org, _admin = tenant
    members = make_members(db, org, 2)
    run = preview(client, members)
    resolve = member_bulk.resolve_org_admin_session
    calls = 0

    def revoked_after_first(*args, **kwargs):
        nonlocal calls
        calls += 1
        return resolve(*args, **kwargs) if calls == 1 else None

    monkeypatch.setattr(member_bulk, "resolve_org_admin_session", revoked_after_first)
    response = client.post(f"{PREFIX}/{run['id']}/execute", json={"request_key": uuid.uuid4().hex})
    assert response.status_code == 401
    assert db.query(EmailOutbox).count() == 0
    assert client.get(f"{PREFIX}/{run['id']}").json()["counts"]["pending"] == 2


def test_two_runs_for_same_recipient_serialize_throttle(client, tenant):
    db, org, _admin = tenant
    member = make_members(db, org)[0]
    first = preview(client, [member])
    second = preview(client, [member])
    keys = [uuid.uuid4().hex, uuid.uuid4().hex]

    def advance(pair):
        run, key = pair
        return client.post(f"{PREFIX}/{run['id']}/execute", json={"request_key": key})

    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(advance, zip((first, second), keys)))
    assert all(response.status_code == 200 for response in responses), [response.text for response in responses]
    outcomes = [response.json()["counts"] for response in responses]
    assert sum(outcome["succeeded"] for outcome in outcomes) == 1
    assert sum(outcome["skipped"] for outcome in outcomes) == 1
    assert db.query(EmailOutbox).count() == 1
