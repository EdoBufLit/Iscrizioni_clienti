from __future__ import annotations

import csv
import io
import uuid
from datetime import datetime, timedelta
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.db import SessionLocal
from app.models import (
    AdminRole,
    AdminUser,
    Member,
    MemberStatus,
    OrgAdminToken,
    Organization,
)
from app.services.forms import export_submissions_csv
from app.utils import hash_token


@pytest.fixture
def db():
    session = SessionLocal()
    yield session
    session.close()


def _login_org_admin(client, db, *, org_id: int, suffix: str) -> None:
    admin = AdminUser(
        email=f"csv-admin-{suffix}@example.com",
        role=AdminRole.ORG_ADMIN,
        org_id=org_id,
        is_active=True,
    )
    db.add(admin)
    db.commit()
    db.refresh(admin)

    raw_token = f"csv-admin-token-{suffix}"
    db.add(
        OrgAdminToken(
            admin_id=admin.id,
            token_hash=hash_token(raw_token),
            expires_at=datetime.utcnow() + timedelta(minutes=15),
        )
    )
    db.commit()

    response = client.get(
        f"/api/org-admin/auth/verify?token={raw_token}",
        follow_redirects=False,
    )
    assert response.status_code == 302, response.text


def test_member_csv_neutralizes_formula_cells_without_changing_normal_values(client, db):
    suffix = uuid.uuid4().hex[:8]
    organization = Organization(
        name=f"CSV Club {suffix}",
        slug=f"csv-club-{suffix}",
        privacy_version="v1",
        is_active=True,
    )
    db.add(organization)
    db.commit()
    db.refresh(organization)
    _login_org_admin(client, db, org_id=organization.id, suffix=suffix)

    member = Member(
        org_id=organization.id,
        first_name="=2+2",
        last_name="\t+SUM(1,1)",
        email='\u200b-HYPERLINK("https://example.test")',
        fiscal_code=f"SAFE{suffix.upper()}",
        phone=" @cmd",
        status=MemberStatus.ACTIVE,
        card_no=900_000 + int(suffix[:5], 16),
        card_year=datetime.utcnow().year,
        joined_at=datetime(2026, 7, 15),
    )
    db.add(member)
    db.commit()

    response = client.get("/api/org-admin/members.csv")

    assert response.status_code == 200, response.text
    assert response.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == [
        "Nome",
        "Cognome",
        "Email",
        "Codice Fiscale",
        "Telefono",
        "Stato",
        "Tessera",
        "Data iscrizione",
    ]
    exported = next(row for row in rows[1:] if row[3] == member.fiscal_code)
    assert exported[0] == "'=2+2"
    assert exported[1] == "'\t+SUM(1,1)"
    assert exported[2] == "'\u200b-HYPERLINK(\"https://example.test\")"
    assert exported[3] == member.fiscal_code
    assert exported[4] == "' @cmd"
    assert exported[7] == "2026-07-15"

    workbook_response = client.get("/api/org-admin/members.xlsx")

    assert workbook_response.status_code == 200, workbook_response.text
    assert (
        workbook_response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(workbook_response.content))
    sheet = workbook["Libro soci"]
    assert sheet.sheet_view.showGridLines is False
    assert sheet.freeze_panes == "A5"
    assert sheet["A1"].value == f"Libro soci - {organization.name}"
    assert sheet["A4"].value == "Nome"
    assert sheet["H4"].value == "Data iscrizione"
    exported_row = next(
        row
        for row in sheet.iter_rows(min_row=5, values_only=False)
        if row[3].value == member.fiscal_code
    )
    assert exported_row[0].value == "'=2+2"
    assert exported_row[4].value == " @cmd"
    assert exported_row[5].value == "Attivo"
    assert exported_row[7].value.date() == datetime(2026, 7, 15).date()
    assert exported_row[7].number_format == "dd/mm/yyyy"
    assert sheet.tables["LibroSociTable"].ref.startswith("A4:H")


def test_form_submission_csv_neutralizes_scalar_and_joined_list_formulas():
    fields = [
        SimpleNamespace(id=1, sort_order=1, field_key="normal_value"),
        SimpleNamespace(id=2, sort_order=2, field_key="equals_formula"),
        SimpleNamespace(id=3, sort_order=3, field_key="spaced_formula"),
        SimpleNamespace(id=4, sort_order=4, field_key="controlled_formula"),
        SimpleNamespace(id=5, sort_order=5, field_key="list_formula"),
    ]
    submission = SimpleNamespace(
        id=17,
        submitted_at=datetime(2026, 7, 15, 12, 30),
        status="pending",
        submitted_by_user_id=None,
        payload_json={
            "normal_value": "Mario Rossi",
            "equals_formula": "=2+2",
            "spaced_formula": "  @SUM(1,1)",
            "controlled_formula": "\u202e-10+20",
            "list_formula": ["+1", "safe"],
        },
    )
    form = SimpleNamespace(fields=fields, submissions=[submission])

    rows = list(csv.DictReader(export_submissions_csv(form)))

    assert len(rows) == 1
    exported = rows[0]
    assert exported["normal_value"] == "Mario Rossi"
    assert exported["equals_formula"] == "'=2+2"
    assert exported["spaced_formula"] == "'  @SUM(1,1)"
    assert exported["controlled_formula"] == "'\u202e-10+20"
    assert exported["list_formula"] == "'+1, safe"
