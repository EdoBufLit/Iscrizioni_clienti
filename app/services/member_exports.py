from __future__ import annotations

from datetime import datetime
from io import BytesIO

from app.models import Member, Organization
from app.services.csv_safety import neutralize_csv_formula


MEMBER_EXPORT_HEADERS = [
    "Nome",
    "Cognome",
    "Email",
    "Codice Fiscale",
    "Telefono",
    "Stato",
    "Tessera",
    "Data iscrizione",
]

MEMBER_STATUS_LABELS = {
    "active": "Attivo",
    "pending_verification": "Verifica in corso",
    "pending_docs": "Documenti da verificare",
    "pending_cards": "Tessera da assegnare",
    "rejected": "Rifiutato",
    "expired": "Scaduto",
}


def member_export_values(member: Member) -> list[object]:
    return [
        neutralize_csv_formula(member.first_name or ""),
        neutralize_csv_formula(member.last_name or ""),
        neutralize_csv_formula(member.email or ""),
        neutralize_csv_formula(member.fiscal_code or ""),
        neutralize_csv_formula(member.phone or ""),
        neutralize_csv_formula(member.status.value if member.status else ""),
        member.card_no if member.card_no is not None else "",
        member.joined_at.date() if member.joined_at else "",
    ]


def _xlsx_text(value: object | None) -> str:
    text = str(value or "")
    # openpyxl treats only a leading "=" as a formula. Other CSV injection
    # prefixes remain ordinary strings in XLSX and should not gain a visible apostrophe.
    return f"'{text}" if text.startswith("=") else text


def member_workbook_values(member: Member) -> list[object]:
    raw_status = member.status.value if member.status else ""
    return [
        _xlsx_text(member.first_name),
        _xlsx_text(member.last_name),
        _xlsx_text(member.email),
        _xlsx_text(member.fiscal_code),
        _xlsx_text(member.phone),
        MEMBER_STATUS_LABELS.get(raw_status, raw_status),
        member.card_no if member.card_no is not None else "",
        member.joined_at.date() if member.joined_at else "",
    ]


def build_members_workbook(
    *,
    organization: Organization,
    members: list[Member],
    generated_at: datetime | None = None,
) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo

    created_at = generated_at or datetime.utcnow()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Libro soci"
    sheet.sheet_view.showGridLines = False

    last_column = "H"
    sheet.merge_cells(f"A1:{last_column}1")
    title_cell = sheet["A1"]
    title_cell.value = f"Libro soci - {organization.name}"
    title_cell.font = Font(name="Aptos Display", size=20, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="5A001F")
    title_cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 38

    sheet.merge_cells(f"A2:{last_column}2")
    subtitle_cell = sheet["A2"]
    subtitle_cell.value = (
        f"Esportato il {created_at.strftime('%d/%m/%Y alle %H:%M')} - "
        f"{len(members)} soci"
    )
    subtitle_cell.font = Font(name="Aptos", size=10, color="5B6472")
    subtitle_cell.fill = PatternFill("solid", fgColor="F7F3F4")
    subtitle_cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[2].height = 24

    header_row = 4
    for column_index, header in enumerate(MEMBER_EXPORT_HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=column_index, value=header)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F4858")
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[header_row].height = 26

    thin_line = Side(style="thin", color="E7E9ED")
    for row_index, member in enumerate(members, start=header_row + 1):
        values = member_workbook_values(member)
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.font = Font(name="Aptos", size=10, color="26313D")
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=column_index in {3, 4},
            )
            cell.border = Border(bottom=thin_line)
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="FAF7F8")
        sheet.cell(row=row_index, column=8).number_format = "dd/mm/yyyy"
        sheet.row_dimensions[row_index].height = 22

    final_row = max(header_row + 1, header_row + len(members))
    if not members:
        for column_index in range(1, len(MEMBER_EXPORT_HEADERS) + 1):
            sheet.cell(row=header_row + 1, column=column_index, value="")

    table = Table(
        displayName="LibroSociTable",
        ref=f"A{header_row}:{last_column}{final_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)

    widths = {
        "A": 20,
        "B": 20,
        "C": 34,
        "D": 22,
        "E": 19,
        "F": 22,
        "G": 14,
        "H": 18,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A{header_row}:{last_column}{final_row}"
    sheet.print_title_rows = f"1:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.5
    sheet.page_margins.bottom = 0.5

    payload = BytesIO()
    workbook.save(payload)
    payload.seek(0)
    return payload
